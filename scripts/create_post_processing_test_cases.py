from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def main():
    out_dir = Path("docs/superpowers/test-cases")
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "2026-06-11-后处理闭环自测用例.xlsx"

    rows = [
        [
            "UT-MEM-001",
            "单元测试",
            "分镜上下文记忆",
            "构建五层记忆并产出 context pack",
            "包含 shot、角色、sourceArtifact",
            "调用 buildStoryboardContext 并设置 currentShotId/tokenBudget",
            "生成 storyboard-context-memory.json；contextPack.currentShotId 正确；保留 pinned/邻近上下文",
            "tests/storyboardContextMemory.test.js",
            "node --test tests\\storyboardContextMemory.test.js",
            "已自动化",
            "覆盖记忆分层、压缩、冲突、artifact",
        ],
        [
            "UT-CVC-001",
            "单元测试",
            "跨视频一致性",
            "HappyHorse provider 才允许写跨视频 memory",
            "projectKey=HappyHorse 且 provider=happyhorse",
            "调用 checkCrossVideoConsistency",
            "contextMemoryPatch.writeAllowed=true",
            "tests/postProcessingLoop.regression.test.js",
            "node --test tests\\postProcessingLoop.regression.test.js",
            "已自动化",
            "新增集中回归",
        ],
        [
            "UT-CVC-002",
            "单元测试",
            "跨视频一致性",
            "Seedance provider 禁止写 HappyHorse memory",
            "projectKey=HappyHorse 但 provider=seedance",
            "调用 checkCrossVideoConsistency",
            "status=block；writeAllowed=false；reason=non_happyhorse_provider",
            "tests/crossVideoConsistency.test.js; tests/postProcessingLoop.regression.test.js",
            "node --test tests\\crossVideoConsistency.test.js tests\\postProcessingLoop.regression.test.js",
            "已自动化",
            "防止误承诺 Seedance 续写能力",
        ],
        [
            "UT-CVC-003",
            "单元测试",
            "跨视频一致性",
            "metadata-only provider 也参与 gate",
            "videoResults 不带 provider，videoMetadata.provider=seedance",
            "调用 checkCrossVideoConsistency",
            "status=block；writeAllowed=false",
            "tests/crossVideoConsistency.test.js; tests/postProcessingLoop.regression.test.js",
            "node --test tests\\crossVideoConsistency.test.js",
            "已自动化",
            "覆盖 subagent review 发现的问题",
        ],
        [
            "UT-AV-001",
            "单元测试",
            "音画包装层",
            "无 BGM/SFX 素材不阻断",
            "assets 为空",
            "调用 buildAvPackagingPlan",
            "出现 audio_assets_missing warning；blocking=false",
            "tests/avPackagingPlan.test.js",
            "node --test tests\\avPackagingPlan.test.js",
            "已自动化",
            "第一版缺素材只 warning",
        ],
        [
            "UT-AV-002",
            "单元测试",
            "音画包装层",
            "配置了不存在的本地素材路径会 warning",
            "assets.bgm/sfx.path 指向不存在文件",
            "调用 buildAvPackagingPlan",
            "cue.assetExists=false；bgm_asset_missing/sfx_asset_missing；blocking=false",
            "tests/avPackagingPlan.test.js; tests/postProcessingLoop.regression.test.js",
            "node --test tests\\avPackagingPlan.test.js tests\\postProcessingLoop.regression.test.js",
            "已自动化",
            "防止静默跳过素材",
        ],
        [
            "UT-AV-003",
            "单元测试",
            "音画包装层",
            "字幕样式传入 videoComposer",
            "packagingPlan.subtitleStyleProfile 存在",
            "调用 composeFromLegacy",
            "ASS 字幕使用包装层字体/字号",
            "tests/videoComposer.test.js",
            "node --test tests\\videoComposer.test.js",
            "已自动化",
            "验证中文字体兼容路径",
        ],
        [
            "UT-COMP-001",
            "单元测试",
            "视频合成",
            "对白镜头优先 lipsync clip",
            "同一 shot 同时有 generated video 和 lipsync clip",
            "调用 buildCompositionPlan",
            "dialogue shot visualType=lipsync_clip",
            "tests/videoComposer.test.js",
            "node --test tests\\videoComposer.test.js",
            "已自动化",
            "防止嘴型被普通视频覆盖",
        ],
        [
            "UT-COMP-002",
            "单元测试",
            "视频合成",
            "本地 BGM/SFX 进入 FFmpeg 混音输入",
            "BGM/SFX 文件存在",
            "调用 collectPackagingAudioItems/buildAudioMixFilter",
            "生成 delayed mix inputs，包含 volume/gain",
            "tests/videoComposer.test.js",
            "node --test tests\\videoComposer.test.js",
            "已自动化",
            "验证包装层真正接入合成",
        ],
        [
            "UT-PCR-001",
            "单元测试",
            "成片预览后编辑闭环",
            "第一版只生成 edit-task-pack，不自动修片",
            "上游 QA 存在风险",
            "调用 buildPostComposeReview/runPostComposeReview",
            "executionMode=manual_only；automaticProviderCalls=false；不修改 final video/compose plan",
            "tests/postComposeReview.test.js",
            "node --test tests\\postComposeReview.test.js",
            "已自动化",
            "核心边界",
        ],
        [
            "UT-PCR-002",
            "单元测试",
            "成片预览后编辑闭环",
            "能消费 composer 原生数组 compose-plan",
            "composePlan 为数组且 visualType=static_image",
            "调用 buildPostComposeReview",
            "生成 replace_clip/regenerate_shot task；带 timelineStartMs/timelineEndMs",
            "tests/postComposeReview.test.js; tests/postProcessingLoop.regression.test.js",
            "node --test tests\\postComposeReview.test.js tests\\postProcessingLoop.regression.test.js",
            "已自动化",
            "覆盖 subagent review 发现的问题",
        ],
        [
            "UT-ART-001",
            "单元测试",
            "运行产物",
            "新增 agent 目录进入 runArtifacts",
            "创建 run artifact context",
            "调用 createRunArtifactContext",
            "存在 storyboardContextAgent、crossVideoConsistency、avPackaging、postComposeReview、humanReviewQueue 目录",
            "tests/runArtifacts.test.js",
            "node --test tests\\runArtifacts.test.js",
            "已自动化",
            "保证每个模块产 artifact",
        ],
        [
            "E2E-LOOP-001",
            "E2E",
            "后处理闭环",
            "四个后处理 agent 串联产出 artifacts",
            "准备临时 run artifact context 和两镜头输入",
            "依次执行 buildStoryboardContext、runCrossVideoConsistency、runAvPackaging、runPostComposeReview、writeHumanReviewQueueArtifacts",
            "产出 memory、cross-video-context-memory、av-packaging-plan、edit-task-pack、human-review-queue；edit task 进入人审队列",
            "tests/postProcessingLoop.e2e.test.js",
            "node --test tests\\postProcessingLoop.e2e.test.js",
            "已自动化",
            "本次新增 e2e",
        ],
        [
            "E2E-DIR-001",
            "E2E",
            "Director 集成",
            "director 跑完整项目链路并接入四模块",
            "使用 stub provider 和临时目录",
            "运行 runEpisodePipeline",
            "packagingPlan 传入 composer；postComposeReview items 合并 humanReviewQueue；最终不自动重跑 provider",
            "tests/director.project-run.test.js",
            "node --test tests\\director.project-run.test.js",
            "已自动化",
            "现有 director 集成覆盖",
        ],
        [
            "ACC-001",
            "验收测试",
            "完整聚焦套件",
            "代码改动完成",
            "运行后处理相关测试全集",
            "执行命令并检查 TAP 结果",
            "全部测试通过，0 fail",
            "多文件",
            "node --test tests\\storyboardContextMemory.test.js tests\\crossVideoConsistency.test.js tests\\avPackagingPlan.test.js tests\\postComposeReview.test.js tests\\postProcessingLoop.regression.test.js tests\\postProcessingLoop.e2e.test.js tests\\runArtifacts.test.js tests\\videoComposer.test.js tests\\director.artifacts.test.js tests\\director.project-run.test.js",
            "已自动化",
            "最终回归命令",
        ],
        [
            "MAN-001",
            "人工验收",
            "产物审查",
            "保留一次真实或 stub run 的 artifact 目录",
            "打开 09o/10/09q/10b/12 目录下 JSON/MD",
            "人工确认报告可读、字段完整、中文不乱码",
            "JSON/Markdown 可读，中文正常，artifact 可追溯",
            "无",
            "人工执行",
            "待人工抽查",
            "建议真实跑一次 HappyHorse-only 项目",
        ],
        [
            "MAN-002",
            "人工验收",
            "Excel 自测矩阵",
            "生成本文件",
            "打开 xlsx 检查筛选、列宽、中文显示",
            "人工检查表格内容",
            "无公式错误；中文可读；用例覆盖四模块",
            str(out_path),
            "人工执行",
            "待人工抽查",
            "本文件无公式",
        ],
    ]

    for row in rows:
        if row[1] == "人工验收":
            row[-1:-1] = ["PENDING_MANUAL", "需要人工打开 artifact 或 Excel 后确认"]
        else:
            row[-1:-1] = ["PASSED", "自动化测试已在 2026-06-11 跑通"]

    headers = [
        "用例ID",
        "测试层级",
        "模块",
        "场景/目标",
        "前置条件",
        "操作步骤",
        "预期结果",
        "自动化文件",
        "运行命令",
        "当前状态",
        "最终验收结果",
        "结果依据",
        "备注",
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "TestCases"
    ws.append(headers)
    for row in rows:
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)

    level_fills = {
        "单元测试": "E2F0D9",
        "E2E": "DDEBF7",
        "验收测试": "FFF2CC",
        "人工验收": "FCE4D6",
    }
    result_fills = {
        "PASSED": "C6EFCE",
        "PENDING_MANUAL": "FFF2CC",
        "FAILED": "FFC7CE",
        "BLOCKED": "D9EAD3",
    }
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
        fill = level_fills.get(row[1].value)
        if fill:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=fill)
        result_cell = row[10]
        result_fill = result_fills.get(result_cell.value)
        if result_fill:
            result_cell.fill = PatternFill("solid", fgColor=result_fill)
            result_cell.font = Font(name="Arial", size=10, bold=True)

    widths = [16, 12, 18, 34, 34, 38, 42, 42, 58, 14, 18, 34, 30]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Summary")
    summary_rows = [
        ["项目", "AI-video-factory-pro 后处理闭环自测"],
        ["日期", "2026-06-11"],
        ["覆盖模块", "分镜上下文记忆；跨视频一致性；音画包装层；成片预览后编辑闭环；人审队列"],
        ["自动化用例数", len([row for row in rows if row[1] != "人工验收"])],
        ["人工验收项", len([row for row in rows if row[1] == "人工验收"])],
        ["PASSED", len([row for row in rows if row[10] == "PASSED"])],
        ["PENDING_MANUAL", len([row for row in rows if row[10] == "PENDING_MANUAL"])],
        ["FAILED", len([row for row in rows if row[10] == "FAILED"])],
        ["核心命令", "node --test tests\\postProcessingLoop.regression.test.js tests\\postProcessingLoop.e2e.test.js"],
    ]
    for row in summary_rows:
        summary.append(row)
    for row in summary.iter_rows():
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    summary["A1"].font = Font(name="Arial", bold=True)
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 120

    legend = wb.create_sheet("Legend")
    legend_rows = [
        ["类型/状态", "颜色", "含义"],
        ["单元测试", "浅绿色整行", "模块级自动化测试，主要验证单个 domain/agent 的确定性行为。"],
        ["E2E", "浅蓝色整行", "端到端链路测试，验证多个后处理 agent 串联和 artifact 产出。"],
        ["验收测试", "浅黄色整行", "聚焦回归套件或发布前验收命令。"],
        ["人工验收", "浅橙色整行", "必须由人工打开 artifact、Excel 或真实运行结果确认。"],
        ["PASSED", "绿色结果单元格", "自动化测试已经通过，本轮验收结果为通过。"],
        ["PENDING_MANUAL", "黄色结果单元格", "自动化无法完全判断，需要人工验收后改为 PASSED 或 FAILED。"],
        ["FAILED", "红色结果单元格", "测试或人工验收失败，需要修复后重新执行。"],
        ["BLOCKED", "浅绿结果单元格", "因外部条件缺失无法执行，例如缺真实 provider、素材或人工输入。"],
    ]
    for row in legend_rows:
        legend.append(row)
    for cell in legend[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
    legend_color_by_label = {
        "单元测试": "E2F0D9",
        "E2E": "DDEBF7",
        "验收测试": "FFF2CC",
        "人工验收": "FCE4D6",
        "PASSED": "C6EFCE",
        "PENDING_MANUAL": "FFF2CC",
        "FAILED": "FFC7CE",
        "BLOCKED": "D9EAD3",
    }
    for row in legend.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
        fill = legend_color_by_label.get(row[0].value)
        if fill:
            row[1].fill = PatternFill("solid", fgColor=fill)
    legend.column_dimensions["A"].width = 20
    legend.column_dimensions["B"].width = 22
    legend.column_dimensions["C"].width = 90

    wb.save(out_path)

    loaded = load_workbook(out_path, data_only=False)
    assert "TestCases" in loaded.sheetnames
    assert "Summary" in loaded.sheetnames
    assert "Legend" in loaded.sheetnames
    assert loaded["TestCases"].max_row == len(rows) + 1
    print(out_path.resolve())


if __name__ == "__main__":
    main()
